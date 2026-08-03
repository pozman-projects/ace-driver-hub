"""EB-12 · Migration Preparation Module.

A read-only migration control layer. Nothing here mutates canonical DCC
records (drivers, owners, vehicles, equipment, relationships, compliance,
activation, documents, number sequences). It only *proposes*, *validates*,
*reconciles* and *previews*.

Collections (all UUID-keyed):
  migration_source_workbooks    ← uploaded/registered spreadsheets
  migration_source_sheets       ← detected sheets within a workbook
  migration_mapping_profiles    ← versioned column→field mappings
  migration_field_mappings      ← rows inside a profile
  migration_transform_rules     ← reusable transformation functions
  migration_source_lineage      ← per-value provenance for completed dry runs
  migration_dry_runs            ← dry-run header
  migration_dry_run_rows        ← one row per source row per dry run
  migration_dry_run_changes     ← proposed field-level changes
  migration_issues              ← validation / reconciliation issues
  migration_reconciliation_results ← aggregated reconciliation reports
  migration_go_no_go_reports    ← final readiness reports
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from openpyxl import load_workbook
from pydantic import BaseModel, Field


# ── Collections ───────────────────────────────────────────────────────────────
WB_COLL = "migration_source_workbooks"
SHEET_COLL = "migration_source_sheets"
PROFILE_COLL = "migration_mapping_profiles"
FIELD_COLL = "migration_field_mappings"
TRANSFORM_COLL = "migration_transform_rules"
LINEAGE_COLL = "migration_source_lineage"
DRYRUN_COLL = "migration_dry_runs"
DRYRUN_ROW_COLL = "migration_dry_run_rows"
DRYRUN_CHG_COLL = "migration_dry_run_changes"
ISSUE_COLL = "migration_issues"
RECON_COLL = "migration_reconciliation_results"
GNG_COLL = "migration_go_no_go_reports"

DRIVERS_COLL = "drivers"
OWNERS_COLL = "owners"
VEHICLES_COLL = "vehicles"
EQUIPMENT_COLL = "equipment"
DOR_COLL = "driver_owner_relationships"
DVA_COLL = "driver_vehicle_assignments"
DEA_COLL = "driver_equipment_assignments"

# ── Controlled enums ──────────────────────────────────────────────────────────
AUTHORITY_LEVELS = {"Authoritative", "Supporting", "Historical", "Reference Only", "Unknown"}
WORKBOOK_STATUSES = {"Uploaded", "Profiled", "Mapping Required", "Mapping Complete",
                     "Validation Failed", "Dry Run Ready", "Archived"}
SHEET_STATUSES = {"Unclassified", "Classified", "Ignored", "Mapping Required",
                  "Mapping Complete", "Validation Failed", "Dry Run Ready"}
PROFILE_STATUSES = {"Draft", "In Review", "Approved", "Superseded", "Archived"}
MAPPING_TYPES = {"Direct", "Transformed", "Lookup", "Reference", "Constant", "Derived", "Ignored"}
NULL_HANDLING = {"Reject Row", "Leave Blank", "Use Default", "Preserve Existing",
                 "Set Not Recorded", "Not Applicable"}
CONFLICT_STRATEGY = {"Reject", "Flag for Review", "Preserve Canonical", "Use Source",
                     "Use Most Recent", "Use Approved Override"}
DRYRUN_STATUSES = {"Queued", "Profiling", "Validating", "Reconciling",
                   "Preview Ready", "Failed", "Archived"}
MATCH_STATUSES = {"No Match", "Exact Match", "Probable Match", "Multiple Matches",
                  "Conflict", "Not Applicable"}
PROPOSED_ACTIONS = {"Create", "Update", "Preserve Existing", "Skip", "Manual Review", "Reject"}
ROW_STATUSES = {"Valid", "Warning", "Blocking", "Ignored"}
CHANGE_TYPES = {"Create Field", "Update Field", "Preserve", "Clear",
                "Relationship Change", "Identifier Change", "Compliance Change", "Document Link"}
ISSUE_SEVERITIES = {"Info", "Warning", "Error", "Critical"}
ISSUE_STATUSES = {"Open", "In Review", "Resolved", "Accepted Risk", "Rejected", "Archived"}
RESOLUTION_TYPES = {"Correct Source Mapping", "Preserve Canonical", "Use Source",
                    "Use Approved Override", "Merge Records", "Create Missing Reference",
                    "Mark Not Applicable", "Reject Row", "Ignore Warning", "Other"}
GNG_RESULTS = {"GO", "CONDITIONAL GO", "NO-GO"}

TARGET_ENTITIES = {"Driver", "Owner", "Vehicle", "Equipment",
                    "DriverOwnerRelationship", "DriverVehicleAssignment",
                    "DriverEquipmentAssignment", "DriverLicence",
                    "VehicleRegistration", "VehicleInsurance", "VehicleInspection",
                    "VehicleDefect", "VehicleMaintenance", "EquipmentCompliance",
                    "Document"}

PROTECTED_FIELDS = {
    "Driver": {"driver_code", "dispatch_number", "abn", "payroll_number"},
    "Owner": {"abn"},
    "Vehicle": {"vin", "registration_number"},
    "Equipment": {"equipment_number"},
}

ROLE_APPROVE = {"Manager", "Admin"}
ROLE_MANAGE = {"Manager", "Admin"}
ROLE_EXECUTE = {"Manager", "Admin", "Compliance"}
ROLE_RESOLVE_ISSUE = {"Manager", "Admin", "Compliance"}
ROLE_VIEW = {"ReadOnly", "Allocator", "Compliance", "Manager", "Admin"}

RESERVED_DISPATCH = {0, 13}


def _uuid() -> str:
    return str(uuid.uuid4())


def _iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _require(user: Dict[str, Any], allowed: set, err="Insufficient permissions"):
    if user.get("role") not in allowed:
        raise HTTPException(status_code=403, detail=err)


def _strip(v: Optional[dict]) -> Optional[dict]:
    if not v:
        return v
    out = dict(v)
    out.pop("_id", None)
    return out


# ── Transform registry (deterministic pure functions) ─────────────────────────
def _t_trim(v, params=None):
    if v is None:
        return None
    return str(v).strip()

def _t_upper(v, params=None):
    return str(v).upper() if v is not None else None

def _t_lower(v, params=None):
    return str(v).lower() if v is not None else None

def _t_title(v, params=None):
    return str(v).title() if v is not None else None

def _t_phone(v, params=None):
    if v is None:
        return None
    s = re.sub(r"[^\d+]", "", str(v))
    return s or None

def _t_email(v, params=None):
    if v is None:
        return None
    s = str(v).strip().lower()
    return s if "@" in s else None

def _t_date(v, params=None):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None

def _t_currency(v, params=None):
    if v is None or v == "":
        return None
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(s)
    except (ValueError, TypeError):
        return None

def _t_percentage(v, params=None):
    if v is None or v == "":
        return None
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        val = float(s)
        return val
    except (ValueError, TypeError):
        return None

def _t_bool(v, params=None):
    if v is None or v == "":
        return None
    s = str(v).strip().lower()
    if s in ("true", "yes", "y", "1", "on"):
        return True
    if s in ("false", "no", "n", "0", "off"):
        return False
    return None

def _t_int(v, params=None):
    if v is None or v == "":
        return None
    try:
        return int(str(v).strip())
    except (ValueError, TypeError):
        return None

def _t_driver_code(v, params=None):
    """Preserve historical non-integer codes as-is; integer-shaped only for live sequence."""
    if v is None or v == "":
        return None
    s = str(v).strip().upper()
    return s

def _t_registration(v, params=None):
    if v is None:
        return None
    return re.sub(r"[^A-Z0-9]", "", str(v).upper())

def _t_vin(v, params=None):
    if v is None:
        return None
    return re.sub(r"[^A-HJ-NPR-Z0-9]", "", str(v).upper())

def _t_abn(v, params=None):
    if v is None:
        return None
    s = re.sub(r"[^\d]", "", str(v))
    return s if len(s) == 11 else None

def _t_status_map(v, params=None):
    m = (params or {}).get("map") or {}
    if v is None:
        return None
    return m.get(str(v).strip(), str(v).strip())

def _t_blank_to_null(v, params=None):
    if v is None:
        return None
    s = str(v).strip()
    return None if s == "" else s

def _t_state_abbr(v, params=None):
    if v is None:
        return None
    m = {"NEW SOUTH WALES": "NSW", "VICTORIA": "VIC", "QUEENSLAND": "QLD",
         "WESTERN AUSTRALIA": "WA", "SOUTH AUSTRALIA": "SA", "TASMANIA": "TAS",
         "AUSTRALIAN CAPITAL TERRITORY": "ACT", "NORTHERN TERRITORY": "NT"}
    s = str(v).strip().upper()
    return m.get(s, s)

def _t_control_chars(v, params=None):
    if v is None:
        return None
    return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", str(v))


TRANSFORM_FUNCTIONS = {
    "trim": _t_trim, "upper": _t_upper, "lower": _t_lower, "title": _t_title,
    "phone": _t_phone, "email": _t_email, "date": _t_date, "currency": _t_currency,
    "percentage": _t_percentage, "bool": _t_bool, "int": _t_int,
    "driver_code": _t_driver_code, "registration": _t_registration, "vin": _t_vin,
    "abn": _t_abn, "status_map": _t_status_map, "blank_to_null": _t_blank_to_null,
    "state_abbr": _t_state_abbr, "control_chars": _t_control_chars,
}


def apply_transform(rule_type: str, value: Any, params: Optional[dict] = None) -> Any:
    fn = TRANSFORM_FUNCTIONS.get(rule_type)
    if not fn:
        return value
    return fn(value, params or {})


# ── Workbook profiling helpers ────────────────────────────────────────────────
def _profile_xlsx(data: bytes) -> List[dict]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets = []
    for idx, name in enumerate(wb.sheetnames):
        ws = wb[name]
        # Detect header row = first non-empty row
        header_row_num = None
        headers: List[str] = []
        first_data_row = None
        row_count = 0
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [(str(c) if c is not None else "") for c in row]
            if not any(v.strip() for v in values):
                continue
            if header_row_num is None:
                header_row_num = r_idx
                headers = [v.strip() for v in values]
                continue
            if first_data_row is None:
                first_data_row = r_idx
            row_count += 1
        col_profile = []
        # Duplicate headers marker
        seen = {}
        for i, h in enumerate(headers):
            if h in seen:
                col_profile.append({"index": i, "name": h, "duplicate_of": seen[h]})
            else:
                seen[h] = i
                col_profile.append({"index": i, "name": h})
        sheets.append({
            "sheet_name": name, "sheet_index": idx,
            "header_row": header_row_num, "first_data_row": first_data_row,
            "detected_column_count": len(headers), "detected_row_count": row_count,
            "column_profile": col_profile,
        })
    return sheets


def _profile_csv(data: bytes) -> List[dict]:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    headers = rows[0] if rows else []
    data_rows = rows[1:] if len(rows) > 1 else []
    col_profile = [{"index": i, "name": h.strip()} for i, h in enumerate(headers)]
    return [{
        "sheet_name": "CSV", "sheet_index": 0,
        "header_row": 1, "first_data_row": 2 if data_rows else None,
        "detected_column_count": len(headers),
        "detected_row_count": len(data_rows), "column_profile": col_profile,
    }]


def _read_xlsx_sheet(data: bytes, sheet_name: str) -> Tuple[List[str], List[List[Any]]]:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[sheet_name]
    headers: List[str] = []
    rows: List[List[Any]] = []
    for r in ws.iter_rows(values_only=True):
        if not headers:
            headers = [(str(c).strip() if c is not None else "") for c in r]
            continue
        rows.append(list(r))
    return headers, rows


# ── Matching engine ───────────────────────────────────────────────────────────
async def _match_driver(db, proposed: Dict[str, Any]) -> Tuple[str, List[dict], List[str]]:
    """Return (match_status, matches, evidence)."""
    evidence = []
    query_or: List[dict] = []
    dcid = proposed.get("id") or proposed.get("driver_id")
    if dcid:
        query_or.append({"id": dcid})
        evidence.append(f"driver_id={dcid}")
    if proposed.get("driver_code"):
        query_or.append({"driver_code": proposed["driver_code"]})
        evidence.append(f"driver_code={proposed['driver_code']}")
    if proposed.get("dispatch_number") not in (None, ""):
        query_or.append({"dispatch_number": proposed["dispatch_number"],
                         "is_archived": {"$ne": True}})
        evidence.append(f"dispatch_number={proposed['dispatch_number']}")
    if proposed.get("email") and proposed.get("mobile_phone"):
        query_or.append({"email": proposed["email"], "mobile_phone": proposed["mobile_phone"]})
        evidence.append("email+mobile")
    if not query_or:
        return "No Match", [], evidence
    docs = await db[DRIVERS_COLL].find({"$or": query_or}, {"_id": 0}).to_list(20)
    if not docs:
        return "No Match", [], evidence
    if len(docs) > 1:
        return "Multiple Matches", docs, evidence
    return "Exact Match", docs, evidence


async def _match_owner(db, proposed: Dict[str, Any]) -> Tuple[str, List[dict], List[str]]:
    evidence = []
    query_or: List[dict] = []
    if proposed.get("id"):
        query_or.append({"id": proposed["id"]})
    if proposed.get("abn"):
        query_or.append({"abn": proposed["abn"]})
        evidence.append(f"abn={proposed['abn']}")
    if proposed.get("name"):
        query_or.append({"name": proposed["name"]})
        evidence.append(f"name={proposed['name']}")
    if not query_or:
        return "No Match", [], evidence
    docs = await db[OWNERS_COLL].find({"$or": query_or}, {"_id": 0}).to_list(20)
    if not docs:
        return "No Match", [], evidence
    if len(docs) == 1:
        return "Exact Match", docs, evidence
    return "Multiple Matches", docs, evidence


async def _match_vehicle(db, proposed: Dict[str, Any]) -> Tuple[str, List[dict], List[str]]:
    evidence = []
    query_or: List[dict] = []
    if proposed.get("vin"):
        query_or.append({"vin": proposed["vin"]})
        evidence.append("vin")
    if proposed.get("registration_number") and proposed.get("state"):
        query_or.append({"registration_number": proposed["registration_number"],
                         "state": proposed["state"]})
        evidence.append("registration+state")
    if not query_or:
        return "No Match", [], evidence
    docs = await db[VEHICLES_COLL].find({"$or": query_or}, {"_id": 0}).to_list(20)
    if not docs:
        return "No Match", [], evidence
    if len(docs) == 1:
        return "Exact Match", docs, evidence
    return "Multiple Matches", docs, evidence


async def _match_equipment(db, proposed: Dict[str, Any]) -> Tuple[str, List[dict], List[str]]:
    evidence = []
    query_or: List[dict] = []
    if proposed.get("equipment_number"):
        query_or.append({"equipment_number": proposed["equipment_number"]})
        evidence.append("equipment_number")
    if not query_or:
        return "No Match", [], evidence
    docs = await db[EQUIPMENT_COLL].find({"$or": query_or}, {"_id": 0}).to_list(20)
    if not docs:
        return "No Match", [], evidence
    if len(docs) == 1:
        return "Exact Match", docs, evidence
    return "Multiple Matches", docs, evidence


# ── Indexes / Seed ────────────────────────────────────────────────────────────
async def ensure_indexes(db):
    for coll, key in [
        (WB_COLL, "migration_source_workbook_id"),
        (SHEET_COLL, "migration_source_sheet_id"),
        (PROFILE_COLL, "migration_mapping_profile_id"),
        (FIELD_COLL, "migration_field_mapping_id"),
        (TRANSFORM_COLL, "migration_transform_rule_id"),
        (DRYRUN_COLL, "migration_dry_run_id"),
        (DRYRUN_ROW_COLL, "migration_dry_run_row_id"),
        (DRYRUN_CHG_COLL, "migration_dry_run_change_id"),
        (ISSUE_COLL, "migration_issue_id"),
        (RECON_COLL, "migration_reconciliation_id"),
        (GNG_COLL, "migration_go_no_go_report_id"),
        (LINEAGE_COLL, "migration_source_lineage_id"),
    ]:
        await db[coll].create_index(key, unique=True)
    await db[SHEET_COLL].create_index("migration_source_workbook_id")
    await db[PROFILE_COLL].create_index("migration_source_workbook_id")
    await db[FIELD_COLL].create_index("migration_mapping_profile_id")
    await db[DRYRUN_ROW_COLL].create_index("migration_dry_run_id")
    await db[DRYRUN_CHG_COLL].create_index("migration_dry_run_row_id")
    await db[ISSUE_COLL].create_index("migration_dry_run_id")


async def seed_transform_rules(db):
    if await db[TRANSFORM_COLL].count_documents({"_source": "seed-eb12"}) > 0:
        return
    defaults = [
        ("trim", "Trim whitespace", "trim", {}),
        ("normalise_email", "Normalise email", "email", {}),
        ("normalise_phone", "Normalise phone", "phone", {}),
        ("parse_date", "Parse date", "date", {}),
        ("parse_boolean", "Parse boolean", "bool", {}),
        ("parse_percentage", "Parse percentage", "percentage", {}),
        ("normalise_driver_code", "Normalise Driver Code", "driver_code", {}),
        ("normalise_registration", "Normalise registration", "registration", {}),
        ("normalise_vin", "Normalise VIN", "vin", {}),
        ("normalise_abn", "Normalise ABN", "abn", {}),
        ("blank_to_null", "Convert blank strings to null", "blank_to_null", {}),
        ("state_abbr", "Normalise state abbreviation", "state_abbr", {}),
        ("control_chars", "Remove invalid control chars", "control_chars", {}),
    ]
    for name, desc, rule_type, params in defaults:
        await db[TRANSFORM_COLL].insert_one({
            "migration_transform_rule_id": _uuid(),
            "name": name, "description": desc, "rule_type": rule_type,
            "parameters": params, "input_type": "string", "output_type": "string",
            "status": "Approved", "created_by": "system-seed",
            "updated_by": "system-seed", "created_at": _iso(),
            "updated_at": _iso(), "is_archived": False, "_source": "seed-eb12",
        })


# ── Pydantic request models ───────────────────────────────────────────────────
class ProfileCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    migration_source_workbook_id: str
    migration_source_sheet_id: Optional[str] = None
    target_entity_type: str


class ProfileUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    target_entity_type: Optional[str] = None
    status: Optional[str] = None


class FieldMappingCreate(BaseModel):
    source_column_name: str
    source_column_index: Optional[int] = 0
    target_field: str
    target_entity_type: str
    mapping_type: str = "Direct"
    required: bool = False
    default_value: Optional[str] = None
    null_handling: str = "Leave Blank"
    transform_rule_ids: List[str] = Field(default_factory=list)
    reference_entity_type: Optional[str] = None
    match_strategy: Optional[str] = None
    conflict_strategy: str = "Flag for Review"
    validation_rule: Optional[str] = None
    display_order: int = 0


class TransformRuleCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    rule_type: str
    parameters: Dict[str, Any] = Field(default_factory=dict)
    input_type: str = "string"
    output_type: str = "string"


class DryRunCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    mapping_profile_ids: List[str]
    source_workbook_ids: List[str]


class IssueResolve(BaseModel):
    resolution_type: str
    resolution_note: Optional[str] = ""


class GoNoGoDecision(BaseModel):
    approval_note: Optional[str] = ""


# ── Router ────────────────────────────────────────────────────────────────────
def build_migration_prep_router(db, get_current_user):
    router = APIRouter(prefix="/api/migration-prep", tags=["migration-prep"])

    # ═══════════════════════════════════════ Workbooks
    @router.post("/workbooks/profile")
    async def profile_workbook(
        file: UploadFile = File(...),
        profile_name: str = Form(...),
        display_name: str = Form(""),
        source_system: str = Form(""),
        business_owner: str = Form(""),
        data_domain: str = Form(""),
        authority_level: str = Form("Unknown"),
        expected_frequency: str = Form(""),
        current=Depends(get_current_user),
    ):
        _require(current, ROLE_MANAGE, "Only Manager or Admin may upload workbooks")
        if authority_level not in AUTHORITY_LEVELS:
            raise HTTPException(status_code=400, detail=f"authority_level must be one of {sorted(AUTHORITY_LEVELS)}")
        data = await file.read()
        if not data:
            raise HTTPException(status_code=400, detail="Empty file")
        original = file.filename or "upload"
        safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", original)
        ext = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
        if ext == "xlsx":
            file_type = "xlsx"
            try:
                sheets = _profile_xlsx(data)
            except Exception as e:  # noqa: BLE001
                raise HTTPException(status_code=400, detail=f"Failed to parse XLSX: {e}")
        elif ext == "csv":
            file_type = "csv"
            sheets = _profile_csv(data)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type .{ext} (allowed: xlsx, csv)")

        wb_id = _uuid()
        checksum = _sha256(data)
        wb = {
            "migration_source_workbook_id": wb_id,
            "profile_name": profile_name, "display_name": display_name or profile_name,
            "original_file_name": original, "sanitised_file_name": safe_name,
            "file_type": file_type, "source_system": source_system,
            "business_owner": business_owner, "data_domain": data_domain,
            "authority_level": authority_level, "expected_frequency": expected_frequency,
            "expected_sheet_count": len(sheets), "detected_sheet_count": len(sheets),
            "file_sha256": checksum, "uploaded_by": current.get("email"),
            "uploaded_at": _iso(), "last_profiled_at": _iso(),
            "status": "Profiled" if sheets else "Validation Failed",
            "notes": "", "created_at": _iso(), "updated_at": _iso(),
            "_source": "runtime", "_data": data,  # raw bytes cached for later reads
        }
        await db[WB_COLL].insert_one(wb)
        for s in sheets:
            await db[SHEET_COLL].insert_one({
                "migration_source_sheet_id": _uuid(),
                "migration_source_workbook_id": wb_id,
                **s, "data_domain": data_domain,
                "target_entity_type": "", "sheet_status": "Unclassified",
                "created_at": _iso(), "updated_at": _iso(),
            })
        return _clean_wb(await db[WB_COLL].find_one({"migration_source_workbook_id": wb_id}, {"_id": 0}))

    def _clean_wb(wb):
        if not wb:
            return wb
        out = dict(wb); out.pop("_data", None); out.pop("_id", None); return out

    @router.get("/workbooks")
    async def list_workbooks(current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[WB_COLL].find({"is_archived": {"$ne": True}}, {"_id": 0}).sort("uploaded_at", -1).to_list(500)
        return [_clean_wb(r) for r in rows]

    @router.get("/workbooks/{wb_id}")
    async def get_workbook(wb_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        wb = await db[WB_COLL].find_one({"migration_source_workbook_id": wb_id}, {"_id": 0})
        if not wb:
            raise HTTPException(status_code=404, detail="Workbook not found")
        return _clean_wb(wb)

    @router.get("/workbooks/{wb_id}/sheets")
    async def list_sheets(wb_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[SHEET_COLL].find({"migration_source_workbook_id": wb_id}, {"_id": 0}).sort("sheet_index", 1).to_list(200)
        return rows

    @router.post("/workbooks/{wb_id}/archive")
    async def archive_workbook(wb_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        r = await db[WB_COLL].update_one(
            {"migration_source_workbook_id": wb_id},
            {"$set": {"is_archived": True, "status": "Archived", "updated_at": _iso()}})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Workbook not found")
        return {"archived": True}

    # ═══════════════════════════════════════ Mapping Profiles
    @router.get("/mapping-profiles")
    async def list_profiles(current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[PROFILE_COLL].find({"is_archived": {"$ne": True}}, {"_id": 0}).sort("updated_at", -1).to_list(500)

    @router.post("/mapping-profiles")
    async def create_profile(payload: ProfileCreate, current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE, "Only Manager or Admin may create mapping profiles")
        if payload.target_entity_type not in TARGET_ENTITIES:
            raise HTTPException(status_code=400, detail=f"target_entity_type must be one of {sorted(TARGET_ENTITIES)}")
        # Determine next version for same workbook+sheet+target
        existing = await db[PROFILE_COLL].find({
            "migration_source_workbook_id": payload.migration_source_workbook_id,
            "migration_source_sheet_id": payload.migration_source_sheet_id,
            "target_entity_type": payload.target_entity_type,
        }, {"profile_version": 1}).to_list(200)
        next_ver = 1 + max([e.get("profile_version", 0) for e in existing] + [0])
        doc = {
            "migration_mapping_profile_id": _uuid(),
            **payload.dict(),
            "profile_version": next_ver, "status": "Draft", "is_active": False,
            "effective_from": None, "effective_to": None,
            "created_by": current.get("email"), "updated_by": current.get("email"),
            "created_at": _iso(), "updated_at": _iso(),
            "is_archived": False, "_source": "runtime",
        }
        await db[PROFILE_COLL].insert_one(doc)
        return _strip(await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": doc["migration_mapping_profile_id"]}, {"_id": 0}))

    @router.get("/mapping-profiles/{profile_id}")
    async def get_profile(profile_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0})
        if not p:
            raise HTTPException(status_code=404, detail="Profile not found")
        return p

    @router.put("/mapping-profiles/{profile_id}")
    async def update_profile(profile_id: str, payload: ProfileUpdate, current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0})
        if not p:
            raise HTTPException(status_code=404, detail="Profile not found")
        if p.get("status") == "Approved":
            raise HTTPException(status_code=400, detail="Approved profiles are immutable — clone to a new version")
        upd = {k: v for k, v in payload.dict(exclude_unset=True).items() if v is not None}
        if "status" in upd and upd["status"] not in PROFILE_STATUSES:
            raise HTTPException(status_code=400, detail="Invalid status")
        if upd.get("status") == "Approved":
            raise HTTPException(status_code=400, detail="Use /approve endpoint to approve")
        upd["updated_at"] = _iso(); upd["updated_by"] = current.get("email")
        await db[PROFILE_COLL].update_one({"migration_mapping_profile_id": profile_id}, {"$set": upd})
        return _strip(await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0}))

    @router.post("/mapping-profiles/{profile_id}/clone")
    async def clone_profile(profile_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0})
        if not p:
            raise HTTPException(status_code=404, detail="Profile not found")
        existing = await db[PROFILE_COLL].find({
            "migration_source_workbook_id": p["migration_source_workbook_id"],
            "migration_source_sheet_id": p.get("migration_source_sheet_id"),
            "target_entity_type": p["target_entity_type"],
        }, {"profile_version": 1}).to_list(200)
        next_ver = 1 + max([e.get("profile_version", 0) for e in existing] + [0])
        new_id = _uuid()
        new_doc = {**p, "migration_mapping_profile_id": new_id,
                   "profile_version": next_ver, "status": "Draft", "is_active": False,
                   "created_at": _iso(), "updated_at": _iso(),
                   "created_by": current.get("email"), "updated_by": current.get("email")}
        new_doc.pop("_id", None)
        await db[PROFILE_COLL].insert_one(new_doc)
        # Clone field mappings
        fms = await db[FIELD_COLL].find({"migration_mapping_profile_id": profile_id}, {"_id": 0}).to_list(500)
        for fm in fms:
            new_fm = {**fm, "migration_field_mapping_id": _uuid(),
                      "migration_mapping_profile_id": new_id,
                      "created_at": _iso(), "updated_at": _iso()}
            new_fm.pop("_id", None)
            await db[FIELD_COLL].insert_one(new_fm)
        return _strip(await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": new_id}, {"_id": 0}))

    @router.post("/mapping-profiles/{profile_id}/approve")
    async def approve_profile(profile_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE, "Only Manager or Admin may approve profiles")
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0})
        if not p:
            raise HTTPException(status_code=404, detail="Profile not found")
        if p["status"] == "Approved":
            return _strip(p)
        # Supersede any active approved profile for same key
        await db[PROFILE_COLL].update_many({
            "migration_source_workbook_id": p["migration_source_workbook_id"],
            "migration_source_sheet_id": p.get("migration_source_sheet_id"),
            "target_entity_type": p["target_entity_type"],
            "is_active": True, "status": "Approved",
            "migration_mapping_profile_id": {"$ne": profile_id},
        }, {"$set": {"status": "Superseded", "is_active": False, "updated_at": _iso()}})
        await db[PROFILE_COLL].update_one(
            {"migration_mapping_profile_id": profile_id},
            {"$set": {"status": "Approved", "is_active": True,
                      "effective_from": _iso(),
                      "updated_by": current.get("email"), "updated_at": _iso()}})
        return _strip(await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0}))

    @router.post("/mapping-profiles/{profile_id}/archive")
    async def archive_profile(profile_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        r = await db[PROFILE_COLL].update_one(
            {"migration_mapping_profile_id": profile_id},
            {"$set": {"is_archived": True, "status": "Archived", "is_active": False, "updated_at": _iso()}})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Profile not found")
        return {"archived": True}

    # ═══════════════════════════════════════ Field mappings + transforms
    @router.get("/mapping-profiles/{profile_id}/fields")
    async def list_fields(profile_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[FIELD_COLL].find(
            {"migration_mapping_profile_id": profile_id, "is_archived": {"$ne": True}},
            {"_id": 0}).sort("display_order", 1).to_list(500)

    @router.post("/mapping-profiles/{profile_id}/fields")
    async def create_field(profile_id: str, payload: FieldMappingCreate,
                            current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0})
        if not p:
            raise HTTPException(status_code=404, detail="Profile not found")
        if p["status"] == "Approved":
            raise HTTPException(status_code=400, detail="Cannot modify Approved profile — clone first")
        if payload.mapping_type not in MAPPING_TYPES:
            raise HTTPException(status_code=400, detail="Invalid mapping_type")
        if payload.null_handling not in NULL_HANDLING:
            raise HTTPException(status_code=400, detail="Invalid null_handling")
        if payload.conflict_strategy not in CONFLICT_STRATEGY:
            raise HTTPException(status_code=400, detail="Invalid conflict_strategy")
        # Guard: Use Source for protected canonical fields
        protected = PROTECTED_FIELDS.get(payload.target_entity_type, set())
        if payload.target_field in protected and payload.conflict_strategy == "Use Source":
            raise HTTPException(status_code=400,
                                 detail=f"'Use Source' not allowed for protected field {payload.target_field}. Use 'Use Approved Override' instead.")
        doc = {
            "migration_field_mapping_id": _uuid(),
            "migration_mapping_profile_id": profile_id,
            **payload.dict(), "created_at": _iso(), "updated_at": _iso(),
            "is_archived": False,
        }
        await db[FIELD_COLL].insert_one(doc)
        return _strip(await db[FIELD_COLL].find_one({"migration_field_mapping_id": doc["migration_field_mapping_id"]}, {"_id": 0}))

    @router.put("/field-mappings/{fm_id}")
    async def update_field(fm_id: str, payload: Dict[str, Any], current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        fm = await db[FIELD_COLL].find_one({"migration_field_mapping_id": fm_id}, {"_id": 0})
        if not fm:
            raise HTTPException(status_code=404, detail="Field mapping not found")
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": fm["migration_mapping_profile_id"]}, {"_id": 0})
        if p and p["status"] == "Approved":
            raise HTTPException(status_code=400, detail="Cannot modify field on Approved profile")
        allowed = {"source_column_name", "source_column_index", "target_field",
                    "mapping_type", "required", "default_value", "null_handling",
                    "transform_rule_ids", "conflict_strategy", "validation_rule",
                    "display_order", "reference_entity_type", "match_strategy"}
        upd = {k: v for k, v in payload.items() if k in allowed}
        upd["updated_at"] = _iso()
        await db[FIELD_COLL].update_one({"migration_field_mapping_id": fm_id}, {"$set": upd})
        return _strip(await db[FIELD_COLL].find_one({"migration_field_mapping_id": fm_id}, {"_id": 0}))

    @router.delete("/field-mappings/{fm_id}")
    async def delete_field(fm_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        fm = await db[FIELD_COLL].find_one({"migration_field_mapping_id": fm_id}, {"_id": 0})
        if not fm:
            raise HTTPException(status_code=404, detail="Field mapping not found")
        p = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": fm["migration_mapping_profile_id"]}, {"_id": 0})
        if p and p["status"] == "Approved":
            raise HTTPException(status_code=400, detail="Cannot delete field on Approved profile")
        await db[FIELD_COLL].delete_one({"migration_field_mapping_id": fm_id})
        return {"deleted": True}

    @router.get("/transform-rules")
    async def list_transforms(current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[TRANSFORM_COLL].find({"is_archived": {"$ne": True}}, {"_id": 0}).to_list(500)

    @router.post("/transform-rules")
    async def create_transform(payload: TransformRuleCreate, current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        if payload.rule_type not in TRANSFORM_FUNCTIONS:
            raise HTTPException(status_code=400, detail=f"Unknown rule_type. Available: {sorted(TRANSFORM_FUNCTIONS.keys())}")
        doc = {
            "migration_transform_rule_id": _uuid(),
            **payload.dict(), "status": "Draft",
            "created_by": current.get("email"), "updated_by": current.get("email"),
            "created_at": _iso(), "updated_at": _iso(), "is_archived": False,
            "_source": "runtime",
        }
        await db[TRANSFORM_COLL].insert_one(doc)
        return _strip(await db[TRANSFORM_COLL].find_one({"migration_transform_rule_id": doc["migration_transform_rule_id"]}, {"_id": 0}))

    @router.put("/transform-rules/{rule_id}")
    async def update_transform(rule_id: str, payload: Dict[str, Any], current=Depends(get_current_user)):
        _require(current, ROLE_MANAGE)
        allowed = {"name", "description", "parameters", "status"}
        upd = {k: v for k, v in payload.items() if k in allowed}
        upd["updated_at"] = _iso()
        r = await db[TRANSFORM_COLL].update_one({"migration_transform_rule_id": rule_id}, {"$set": upd})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Transform not found")
        return _strip(await db[TRANSFORM_COLL].find_one({"migration_transform_rule_id": rule_id}, {"_id": 0}))

    # ═══════════════════════════════════════ Dry runs
    @router.post("/dry-runs")
    async def create_dry_run(payload: DryRunCreate, current=Depends(get_current_user)):
        _require(current, ROLE_EXECUTE)
        doc = {
            "migration_dry_run_id": _uuid(), **payload.dict(),
            "status": "Queued", "requested_by": current.get("email"),
            "requested_at": _iso(), "started_at": None, "completed_at": None,
            "failed_at": None, "failure_reason": None,
            "row_count": 0, "valid_row_count": 0, "warning_row_count": 0,
            "invalid_row_count": 0, "proposed_create_count": 0,
            "proposed_update_count": 0, "proposed_unchanged_count": 0,
            "proposed_skip_count": 0, "blocking_issue_count": 0,
            "warning_issue_count": 0, "correlation_id": _uuid(),
            "created_at": _iso(), "updated_at": _iso(), "_source": "runtime",
        }
        await db[DRYRUN_COLL].insert_one(doc)
        return _strip(await db[DRYRUN_COLL].find_one({"migration_dry_run_id": doc["migration_dry_run_id"]}, {"_id": 0}))

    @router.get("/dry-runs")
    async def list_dry_runs(current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[DRYRUN_COLL].find({}, {"_id": 0}).sort("created_at", -1).to_list(200)

    @router.get("/dry-runs/{dr_id}")
    async def get_dry_run(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        d = await db[DRYRUN_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0})
        if not d:
            raise HTTPException(status_code=404, detail="Dry run not found")
        return d

    async def _execute_dry_run(dr_id: str, actor_email: str):
        dr = await db[DRYRUN_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0})
        if not dr:
            raise HTTPException(status_code=404, detail="Dry run not found")
        # Clear any previous rows/changes/issues from earlier runs of this dry run
        await db[DRYRUN_ROW_COLL].delete_many({"migration_dry_run_id": dr_id})
        await db[DRYRUN_CHG_COLL].delete_many({"migration_dry_run_id": dr_id})
        await db[ISSUE_COLL].delete_many({"migration_dry_run_id": dr_id})
        await db[LINEAGE_COLL].delete_many({"migration_dry_run_id": dr_id})

        await db[DRYRUN_COLL].update_one({"migration_dry_run_id": dr_id},
                                         {"$set": {"status": "Validating", "started_at": _iso()}})
        totals = {"row_count": 0, "valid_row_count": 0, "warning_row_count": 0,
                  "invalid_row_count": 0, "proposed_create_count": 0,
                  "proposed_update_count": 0, "proposed_unchanged_count": 0,
                  "proposed_skip_count": 0, "blocking_issue_count": 0,
                  "warning_issue_count": 0}
        # Track duplicates within source (per profile)
        try:
            for profile_id in dr["mapping_profile_ids"]:
                profile = await db[PROFILE_COLL].find_one({"migration_mapping_profile_id": profile_id}, {"_id": 0})
                if not profile:
                    continue
                wb = await db[WB_COLL].find_one({"migration_source_workbook_id": profile["migration_source_workbook_id"]})
                if not wb:
                    continue
                data = wb.get("_data")
                if not data:
                    continue
                sheet = None
                if profile.get("migration_source_sheet_id"):
                    sheet = await db[SHEET_COLL].find_one({"migration_source_sheet_id": profile["migration_source_sheet_id"]}, {"_id": 0})
                sheet_name = sheet["sheet_name"] if sheet else ("CSV" if wb["file_type"] == "csv" else "Sheet1")
                # Read data
                if wb["file_type"] == "xlsx":
                    headers, rows = _read_xlsx_sheet(data, sheet_name)
                else:
                    reader = csv.reader(io.StringIO(data.decode("utf-8", errors="replace")))
                    all_rows = list(reader)
                    headers = all_rows[0] if all_rows else []
                    rows = all_rows[1:]

                fields = await db[FIELD_COLL].find({"migration_mapping_profile_id": profile_id,
                                                     "is_archived": {"$ne": True}}, {"_id": 0}).to_list(500)
                target_type = profile["target_entity_type"]

                # Build column-index map (source_column_name -> index) from headers
                header_idx = {h.strip(): i for i, h in enumerate(headers) if h}
                # Track seen keys per row for internal duplicate detection
                seen_keys: Dict[str, Dict[str, int]] = {"driver_code": {}, "dispatch_number": {},
                                                          "vin": {}, "abn": {}, "email": {}}

                for row_num, row in enumerate(rows, start=2):
                    row_id = _uuid()
                    source_snapshot = {}
                    transformed = {}
                    lineage_entries: List[dict] = []
                    row_issues: List[dict] = []
                    row_status = "Valid"

                    for f in fields:
                        col_name = f["source_column_name"]
                        col_i = header_idx.get(col_name, f.get("source_column_index", 0))
                        raw = row[col_i] if col_i < len(row) else None
                        source_snapshot[col_name] = raw
                        val = raw
                        # Apply transform rules
                        for rid in f.get("transform_rule_ids") or []:
                            rule = await db[TRANSFORM_COLL].find_one({"migration_transform_rule_id": rid})
                            if rule:
                                val = apply_transform(rule["rule_type"], val, rule.get("parameters"))
                        # Null handling
                        if (val is None or val == "") and f.get("required"):
                            if f.get("null_handling") == "Use Default":
                                val = f.get("default_value")
                            elif f.get("null_handling") == "Reject Row":
                                row_status = "Blocking"
                                row_issues.append(_mk_issue(dr_id, row_id, "REQUIRED_FIELD_MISSING",
                                                              "Error", target_type, f["target_field"],
                                                              f"Required field '{f['target_field']}' missing",
                                                              raw, val, True))
                        transformed[f["target_field"]] = val
                        lineage_entries.append({
                            "migration_source_lineage_id": _uuid(),
                            "migration_dry_run_id": dr_id,
                            "migration_source_workbook_id": wb["migration_source_workbook_id"],
                            "workbook_file_name": wb["sanitised_file_name"],
                            "workbook_checksum": wb["file_sha256"],
                            "migration_source_sheet_id": sheet["migration_source_sheet_id"] if sheet else None,
                            "sheet_name": sheet_name, "source_row_number": row_num,
                            "source_column_name": col_name, "source_column_index": col_i,
                            "original_source_value": (str(raw) if raw is not None else None),
                            "transformed_value": (str(val) if val is not None else None),
                            "applied_transform_rules": f.get("transform_rule_ids") or [],
                            "migration_mapping_profile_id": profile_id,
                            "profile_version": profile["profile_version"],
                            "target_entity_type": target_type,
                            "proposed_target_field": f["target_field"],
                            "created_at": _iso(),
                        })

                    # Match
                    if target_type == "Driver":
                        status, matches, ev = await _match_driver(db, transformed)
                    elif target_type == "Owner":
                        status, matches, ev = await _match_owner(db, transformed)
                    elif target_type == "Vehicle":
                        status, matches, ev = await _match_vehicle(db, transformed)
                    elif target_type == "Equipment":
                        status, matches, ev = await _match_equipment(db, transformed)
                    else:
                        status, matches, ev = "Not Applicable", [], []

                    proposed_action = "Create"
                    target_id = None
                    if status == "Exact Match" and matches:
                        proposed_action = "Update"
                        target_id = matches[0].get("id")
                    elif status == "Multiple Matches":
                        proposed_action = "Manual Review"
                        row_status = "Blocking"
                        row_issues.append(_mk_issue(dr_id, row_id, "MULTIPLE_MATCHES",
                                                      "Error", target_type, None,
                                                      f"Multiple {target_type} records matched — manual resolution required",
                                                      None, None, True,
                                                      related=[m.get("id") for m in matches]))
                    elif status == "Probable Match":
                        proposed_action = "Manual Review"

                    # Internal duplicate detection
                    for key in ("driver_code", "dispatch_number", "vin", "abn", "email"):
                        val = transformed.get(key)
                        if val is not None and val != "":
                            if val in seen_keys[key]:
                                row_status = "Blocking"
                                row_issues.append(_mk_issue(dr_id, row_id, f"DUPLICATE_{key.upper()}",
                                                              "Error", target_type, key,
                                                              f"Duplicate {key}={val} within source workbook (row {seen_keys[key][val]})",
                                                              val, val, True))
                            else:
                                seen_keys[key][val] = row_num

                    # Reserved dispatch numbers
                    dnum = transformed.get("dispatch_number")
                    if dnum is not None:
                        try:
                            dnum_int = int(dnum)
                            if dnum_int in RESERVED_DISPATCH:
                                row_status = "Blocking"
                                row_issues.append(_mk_issue(dr_id, row_id, "RESERVED_DISPATCH",
                                                              "Error", target_type, "dispatch_number",
                                                              f"Dispatch number {dnum_int} is reserved",
                                                              dnum, dnum, True))
                        except (ValueError, TypeError):
                            pass

                    # Driver Code integer-shape check
                    dcode = transformed.get("driver_code")
                    if dcode is not None and target_type == "Driver":
                        try:
                            int(str(dcode))
                        except (ValueError, TypeError):
                            row_issues.append(_mk_issue(dr_id, row_id, "NON_INTEGER_DRIVER_CODE",
                                                          "Warning", target_type, "driver_code",
                                                          f"Non-integer Driver Code '{dcode}' — historical only, will not advance live sequence",
                                                          dcode, dcode, False))
                            if row_status == "Valid":
                                row_status = "Warning"

                    # Write issues to DB
                    for iss in row_issues:
                        await db[ISSUE_COLL].insert_one(iss)

                    # Write row
                    await db[DRYRUN_ROW_COLL].insert_one({
                        "migration_dry_run_row_id": row_id,
                        "migration_dry_run_id": dr_id,
                        "source_workbook_id": wb["migration_source_workbook_id"],
                        "source_sheet_id": sheet["migration_source_sheet_id"] if sheet else None,
                        "source_row_number": row_num,
                        "target_entity_type": target_type,
                        "match_status": status,
                        "proposed_action": proposed_action,
                        "target_entity_id": target_id,
                        "row_status": row_status,
                        "source_snapshot": source_snapshot,
                        "transformed_snapshot": transformed,
                        "proposed_target_snapshot": transformed,
                        "issue_count": len(row_issues),
                        "match_evidence": ev,
                        "created_at": _iso(),
                    })

                    # Write proposed changes
                    current_doc = matches[0] if matches else {}
                    for target_field, new_val in transformed.items():
                        cur_val = current_doc.get(target_field)
                        if cur_val == new_val:
                            ctype = "Preserve"
                        elif cur_val is None:
                            ctype = "Create Field"
                        else:
                            ctype = "Update Field"
                        await db[DRYRUN_CHG_COLL].insert_one({
                            "migration_dry_run_change_id": _uuid(),
                            "migration_dry_run_row_id": row_id,
                            "migration_dry_run_id": dr_id,
                            "target_entity_type": target_type,
                            "target_entity_id": target_id,
                            "target_field": target_field,
                            "current_value": cur_val,
                            "proposed_value": new_val,
                            "change_type": ctype,
                            "conflict_status": ("Conflict" if cur_val is not None and cur_val != new_val else "None"),
                            "resolution_status": "None",
                            "resolution_value": None, "resolved_by": None,
                            "resolved_at": None, "created_at": _iso(),
                        })

                    # Lineage
                    for l in lineage_entries:
                        await db[LINEAGE_COLL].insert_one(l)

                    totals["row_count"] += 1
                    if row_status == "Blocking":
                        totals["invalid_row_count"] += 1
                        totals["blocking_issue_count"] += sum(1 for i in row_issues if i["blocking"])
                    elif row_status == "Warning":
                        totals["warning_row_count"] += 1
                        totals["warning_issue_count"] += sum(1 for i in row_issues if not i["blocking"])
                    else:
                        totals["valid_row_count"] += 1
                    if proposed_action == "Create":
                        totals["proposed_create_count"] += 1
                    elif proposed_action == "Update":
                        totals["proposed_update_count"] += 1
                    elif proposed_action == "Manual Review":
                        totals["proposed_skip_count"] += 1

            await db[DRYRUN_COLL].update_one({"migration_dry_run_id": dr_id},
                                              {"$set": {"status": "Preview Ready",
                                                        "completed_at": _iso(),
                                                        **totals, "updated_at": _iso()}})
        except Exception as e:  # noqa: BLE001
            await db[DRYRUN_COLL].update_one({"migration_dry_run_id": dr_id},
                                              {"$set": {"status": "Failed",
                                                        "failed_at": _iso(),
                                                        "failure_reason": str(e),
                                                        "updated_at": _iso()}})
            raise
        return await db[DRYRUN_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0})

    def _mk_issue(dr_id, row_id, code, severity, entity, field, msg, src, prop, blocking, related=None):
        return {
            "migration_issue_id": _uuid(),
            "migration_dry_run_id": dr_id,
            "migration_dry_run_row_id": row_id,
            "issue_code": code, "severity": severity, "category": entity,
            "entity_type": entity, "field_name": field, "message": msg,
            "source_value": src, "proposed_value": prop,
            "related_entity_ids": related or [], "blocking": blocking,
            "status": "Open", "resolution_type": None, "resolution_note": "",
            "resolved_by": None, "resolved_at": None,
            "created_at": _iso(), "updated_at": _iso(),
        }

    @router.post("/dry-runs/{dr_id}/execute")
    async def execute_dry_run(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_EXECUTE)
        return await _execute_dry_run(dr_id, current.get("email"))

    @router.post("/dry-runs/{dr_id}/rerun")
    async def rerun_dry_run(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_EXECUTE)
        return await _execute_dry_run(dr_id, current.get("email"))

    @router.post("/dry-runs/{dr_id}/archive")
    async def archive_dry_run(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        r = await db[DRYRUN_COLL].update_one({"migration_dry_run_id": dr_id},
                                              {"$set": {"status": "Archived", "updated_at": _iso()}})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Dry run not found")
        return {"archived": True}

    @router.get("/dry-runs/{dr_id}/rows")
    async def list_rows(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id}, {"_id": 0}).sort("source_row_number", 1).to_list(2000)

    @router.get("/dry-runs/{dr_id}/changes")
    async def list_changes(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[DRYRUN_CHG_COLL].find({"migration_dry_run_id": dr_id}, {"_id": 0}).to_list(5000)

    # ═══════════════════════════════════════ Issues
    @router.get("/dry-runs/{dr_id}/issues")
    async def list_issues(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        return await db[ISSUE_COLL].find({"migration_dry_run_id": dr_id}, {"_id": 0}).sort([("severity", -1), ("created_at", 1)]).to_list(2000)

    @router.get("/issues/{issue_id}")
    async def get_issue(issue_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        d = await db[ISSUE_COLL].find_one({"migration_issue_id": issue_id}, {"_id": 0})
        if not d:
            raise HTTPException(status_code=404, detail="Issue not found")
        return d

    @router.post("/issues/{issue_id}/resolve")
    async def resolve_issue(issue_id: str, payload: IssueResolve, current=Depends(get_current_user)):
        _require(current, ROLE_RESOLVE_ISSUE)
        if payload.resolution_type not in RESOLUTION_TYPES:
            raise HTTPException(status_code=400, detail="Invalid resolution_type")
        r = await db[ISSUE_COLL].update_one({"migration_issue_id": issue_id},
                                             {"$set": {"status": "Resolved",
                                                        "resolution_type": payload.resolution_type,
                                                        "resolution_note": payload.resolution_note,
                                                        "resolved_by": current.get("email"),
                                                        "resolved_at": _iso(),
                                                        "updated_at": _iso()}})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Issue not found")
        return _strip(await db[ISSUE_COLL].find_one({"migration_issue_id": issue_id}, {"_id": 0}))

    @router.post("/issues/{issue_id}/reopen")
    async def reopen_issue(issue_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_RESOLVE_ISSUE)
        r = await db[ISSUE_COLL].update_one({"migration_issue_id": issue_id},
                                             {"$set": {"status": "Open",
                                                        "resolution_type": None,
                                                        "resolved_by": None,
                                                        "resolved_at": None,
                                                        "updated_at": _iso()}})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Issue not found")
        return _strip(await db[ISSUE_COLL].find_one({"migration_issue_id": issue_id}, {"_id": 0}))

    # ═══════════════════════════════════════ Reconciliation reports
    @router.get("/dry-runs/{dr_id}/reconciliation")
    async def reconciliation(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id}, {"_id": 0}).to_list(5000)
        by_entity = {}
        by_match = {}
        for r in rows:
            by_entity[r["target_entity_type"]] = by_entity.get(r["target_entity_type"], 0) + 1
            by_match[r["match_status"]] = by_match.get(r["match_status"], 0) + 1
        # Duplicate groups
        dup_issues = await db[ISSUE_COLL].find({"migration_dry_run_id": dr_id,
                                                  "issue_code": {"$regex": "^DUPLICATE_"}},
                                                 {"_id": 0}).to_list(500)
        return {"by_entity": by_entity, "by_match_status": by_match,
                "duplicate_issues": dup_issues, "row_total": len(rows)}

    @router.get("/dry-runs/{dr_id}/identifier-impact")
    async def identifier_impact(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id,
                                                 "target_entity_type": "Driver"}, {"_id": 0}).to_list(5000)
        valid_codes, invalid_codes, dup_codes = [], [], {}
        dispatch_values = []
        reserved = []
        historical = []
        for r in rows:
            t = r.get("transformed_snapshot") or {}
            code = t.get("driver_code")
            if code is not None:
                try:
                    int(str(code))
                    valid_codes.append(code)
                except (ValueError, TypeError):
                    invalid_codes.append(code)
                    historical.append(code)
                dup_codes[code] = dup_codes.get(code, 0) + 1
            dnum = t.get("dispatch_number")
            if dnum is not None:
                try:
                    n = int(dnum)
                    if n in RESERVED_DISPATCH:
                        reserved.append(n)
                    else:
                        dispatch_values.append(n)
                except (ValueError, TypeError):
                    pass
        duplicates = {k: v for k, v in dup_codes.items() if v > 1}
        # Projected next sequence (peek only)
        seq = await db["number_sequences"].find_one({"sequence_key": "driver_code_live"}, {"_id": 0})
        current_live = (seq or {}).get("last_value", 0) if seq else 0
        return {
            "driver_code": {
                "valid": len(valid_codes), "invalid": len(invalid_codes),
                "duplicates": duplicates, "historical_only": historical,
                "projected_next_live": current_live + 1,
                "notes": "Dry run does not advance live sequence.",
            },
            "dispatch_number": {
                "reserved_hits": reserved,
                "active_proposals": dispatch_values,
                "reserved_permanent": sorted(RESERVED_DISPATCH),
                "notes": "0 and 13 are permanently reserved. Dry run does not consume reservations.",
            },
        }

    @router.get("/dry-runs/{dr_id}/relationship-impact")
    async def relationship_impact(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id,
                                                 "target_entity_type": {"$in":
                                                    ["DriverOwnerRelationship",
                                                     "DriverVehicleAssignment",
                                                     "DriverEquipmentAssignment"]}}, {"_id": 0}).to_list(5000)
        return {"proposed_relationships": rows,
                 "notes": "No canonical relationships are mutated. Preview only."}

    @router.get("/dry-runs/{dr_id}/compliance-impact")
    async def compliance_impact(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id,
                                                 "target_entity_type": {"$in":
                                                    ["DriverLicence", "VehicleRegistration",
                                                     "VehicleInsurance", "VehicleInspection",
                                                     "VehicleDefect", "VehicleMaintenance",
                                                     "EquipmentCompliance"]}}, {"_id": 0}).to_list(5000)
        return {"proposed_compliance": rows,
                 "notes": "Compliance records are NOT changed. Worst-Status-Wins figures are inferred, not applied."}

    @router.get("/dry-runs/{dr_id}/activation-impact")
    async def activation_impact(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        drivers = await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id,
                                                    "target_entity_type": "Driver"}, {"_id": 0}).to_list(5000)
        preview = []
        for d in drivers:
            t = d.get("transformed_snapshot") or {}
            missing = []
            if not t.get("driver_code"): missing.append("Driver Code")
            if not t.get("dispatch_number"): missing.append("Dispatch Number")
            if not t.get("mobile_phone"): missing.append("Mobile phone")
            projected = "Ready" if not missing else "Blocked"
            preview.append({
                "source_row_number": d["source_row_number"],
                "target_entity_id": d.get("target_entity_id"),
                "projected_activation_status": projected,
                "projected_readiness_status": projected,
                "projected_blocking_items": missing,
                "notes": "Preview only — no checklist, events, or notifications generated.",
            })
        return {"preview": preview}

    @router.get("/dry-runs/{dr_id}/document-manifest")
    async def document_manifest(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        rows = await db[DRYRUN_ROW_COLL].find({"migration_dry_run_id": dr_id,
                                                 "target_entity_type": "Document"}, {"_id": 0}).to_list(5000)
        return {"proposed_documents": rows,
                 "notes": "Manifest-only. Real files are NOT uploaded in EB-12."}

    @router.get("/dry-runs/{dr_id}/rollback-manifest")
    async def rollback_manifest(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        changes = await db[DRYRUN_CHG_COLL].find({"migration_dry_run_id": dr_id}, {"_id": 0}).to_list(5000)
        creates = [c for c in changes if c["change_type"] == "Create Field"]
        updates = [c for c in changes if c["change_type"] == "Update Field"]
        return {
            "would_create_field_count": len(creates),
            "would_update_field_count": len(updates),
            "prior_values": [{"field": c["target_field"],
                               "entity_type": c["target_entity_type"],
                               "entity_id": c.get("target_entity_id"),
                               "current_value": c.get("current_value")} for c in updates],
            "deterministic": True,
            "applied": False,
            "notes": "This is a rollback PREVIEW. EB-12 does not apply any changes.",
        }

    # ═══════════════════════════════════════ Go / No-Go
    @router.post("/dry-runs/{dr_id}/go-no-go")
    async def compute_go_no_go(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        dr = await db[DRYRUN_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0})
        if not dr:
            raise HTTPException(status_code=404, detail="Dry run not found")
        open_blockers = await db[ISSUE_COLL].count_documents({"migration_dry_run_id": dr_id,
                                                                "status": "Open", "blocking": True})
        critical = await db[ISSUE_COLL].count_documents({"migration_dry_run_id": dr_id,
                                                           "severity": "Critical", "status": {"$ne": "Resolved"}})
        warnings = await db[ISSUE_COLL].count_documents({"migration_dry_run_id": dr_id,
                                                           "severity": "Warning", "status": "Open"})
        # Require approved profiles
        profile_ids = dr["mapping_profile_ids"]
        approved = await db[PROFILE_COLL].count_documents({"migration_mapping_profile_id": {"$in": profile_ids},
                                                             "status": "Approved"})
        all_approved = approved == len(profile_ids) and approved > 0
        if open_blockers > 0 or critical > 0 or not all_approved:
            result = "NO-GO"
        elif warnings > 0:
            result = "CONDITIONAL GO"
        else:
            result = "GO"
        # Upsert one report per dry run
        rid = _uuid()
        existing = await db[GNG_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0})
        report = {
            "migration_go_no_go_report_id": (existing or {}).get("migration_go_no_go_report_id") or rid,
            "migration_dry_run_id": dr_id,
            "result": result,
            "computed_by": current.get("email"),
            "computed_at": _iso(),
            "open_blocking_issue_count": open_blockers,
            "critical_issue_count": critical,
            "warning_issue_count": warnings,
            "approved_profiles_ok": all_approved,
            "rollback_ready": True,
            "approval_status": "Pending",
            "approved_by": None, "approved_at": None,
            "rejected_by": None, "rejected_at": None,
            "notes": "",
            "created_at": (existing or {}).get("created_at") or _iso(),
            "updated_at": _iso(),
        }
        if existing:
            await db[GNG_COLL].update_one({"migration_dry_run_id": dr_id}, {"$set": report})
        else:
            await db[GNG_COLL].insert_one(report)
        return _strip(await db[GNG_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0}))

    @router.get("/dry-runs/{dr_id}/go-no-go")
    async def get_go_no_go(dr_id: str, current=Depends(get_current_user)):
        _require(current, ROLE_VIEW)
        r = await db[GNG_COLL].find_one({"migration_dry_run_id": dr_id}, {"_id": 0})
        if not r:
            raise HTTPException(status_code=404, detail="No Go/No-Go report yet")
        return r

    @router.post("/go-no-go/{report_id}/approve")
    async def approve_gng(report_id: str, payload: GoNoGoDecision, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        r = await db[GNG_COLL].find_one({"migration_go_no_go_report_id": report_id}, {"_id": 0})
        if not r:
            raise HTTPException(status_code=404, detail="Report not found")
        if r["result"] == "NO-GO":
            raise HTTPException(status_code=400, detail="Cannot approve a NO-GO report")
        await db[GNG_COLL].update_one({"migration_go_no_go_report_id": report_id},
                                       {"$set": {"approval_status": "Approved",
                                                  "approved_by": current.get("email"),
                                                  "approved_at": _iso(),
                                                  "notes": payload.approval_note,
                                                  "updated_at": _iso()}})
        return _strip(await db[GNG_COLL].find_one({"migration_go_no_go_report_id": report_id}, {"_id": 0}))

    @router.post("/go-no-go/{report_id}/reject")
    async def reject_gng(report_id: str, payload: GoNoGoDecision, current=Depends(get_current_user)):
        _require(current, ROLE_APPROVE)
        r = await db[GNG_COLL].update_one({"migration_go_no_go_report_id": report_id},
                                            {"$set": {"approval_status": "Rejected",
                                                       "rejected_by": current.get("email"),
                                                       "rejected_at": _iso(),
                                                       "notes": payload.approval_note,
                                                       "updated_at": _iso()}})
        if not r.matched_count:
            raise HTTPException(status_code=404, detail="Report not found")
        return _strip(await db[GNG_COLL].find_one({"migration_go_no_go_report_id": report_id}, {"_id": 0}))

    return router
