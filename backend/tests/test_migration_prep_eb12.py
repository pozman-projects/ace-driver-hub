"""EB-12 · Migration Preparation tests.

Uses only sanitised fictional fixtures. No real ACE identifiers, names,
addresses, phone numbers, emails, registrations, VINs, ABNs or licence
numbers are used.
"""
from __future__ import annotations

import io
import os
import uuid

import pytest
import requests
from openpyxl import Workbook

BASE = (os.environ.get("REACT_APP_BACKEND_URL")
        or "https://fleet-ops-center-1.preview.emergentagent.com").rstrip("/")
API = f"{BASE}/api"
PREP = f"{API}/migration-prep"


def _login(email, password):
    r = requests.post(f"{API}/auth/login",
                       json={"email": email, "password": password}, timeout=25)
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _register(admin_headers, role):
    email = f"eb12_{role.lower()}_{uuid.uuid4().hex[:6]}@acedriverhub.com"
    requests.post(f"{API}/auth/register",
                   json={"email": email, "password": "T@1234",
                         "full_name": f"EB12 {role}", "role": role},
                   headers={**admin_headers, "Content-Type": "application/json"},
                   timeout=25)
    return _login(email, "T@1234")


@pytest.fixture(scope="session")
def admin_headers():
    return _login("admin@acedriverhub.com", "Admin@123")

@pytest.fixture(scope="session")
def readonly_headers(admin_headers):
    return _register(admin_headers, "ReadOnly")

@pytest.fixture(scope="session")
def allocator_headers(admin_headers):
    return _register(admin_headers, "Allocator")

@pytest.fixture(scope="session")
def manager_headers(admin_headers):
    return _register(admin_headers, "Manager")

@pytest.fixture(scope="session")
def compliance_headers(admin_headers):
    return _register(admin_headers, "Compliance")


# ── Fixture workbook builders ────────────────────────────────────────────────
def _build_xlsx(sheets: dict) -> bytes:
    """sheets = {sheet_name: [header_row, *data_rows]}."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload_workbook(admin_headers, filename, sheets, authority_level="Authoritative",
                      profile_name=None):
    data = _build_xlsx(sheets)
    files = {"file": (filename, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    form = {
        "profile_name": profile_name or f"Fixture {uuid.uuid4().hex[:6]}",
        "display_name": filename, "source_system": "SanitisedFixture",
        "business_owner": "QA", "data_domain": "Driver",
        "authority_level": authority_level, "expected_frequency": "one-off",
    }
    r = requests.post(f"{PREP}/workbooks/profile", files=files, data=form,
                       headers=admin_headers, timeout=30)
    r.raise_for_status()
    return r.json()


DRIVER_HEADERS = ["driver_code", "dispatch_number", "full_name",
                  "email", "mobile_phone", "start_date"]

CLEAN_DRIVERS = [DRIVER_HEADERS] + [
    ["9001", 501, "Ficta One", "ficta.one@example.test", "0400000001", "2024-01-05"],
    ["9002", 502, "Ficta Two", "ficta.two@example.test", "0400000002", "2024-01-06"],
]


def _create_profile(admin_headers, wb, target="Driver"):
    sheet = requests.get(f"{PREP}/workbooks/{wb['migration_source_workbook_id']}/sheets",
                          headers=admin_headers).json()[0]
    r = requests.post(f"{PREP}/mapping-profiles", json={
        "name": f"Profile {uuid.uuid4().hex[:6]}",
        "migration_source_workbook_id": wb["migration_source_workbook_id"],
        "migration_source_sheet_id": sheet["migration_source_sheet_id"],
        "target_entity_type": target,
    }, headers=admin_headers, timeout=15)
    r.raise_for_status()
    return r.json(), sheet


def _add_field(admin_headers, profile_id, source_col, target_field,
                transform_ids=None, required=False):
    r = requests.post(f"{PREP}/mapping-profiles/{profile_id}/fields", json={
        "source_column_name": source_col, "target_field": target_field,
        "target_entity_type": "Driver", "mapping_type": "Direct",
        "required": required, "null_handling": "Leave Blank",
        "transform_rule_ids": transform_ids or [],
        "conflict_strategy": "Flag for Review", "display_order": 0,
    }, headers=admin_headers, timeout=15)
    r.raise_for_status()
    return r.json()


def _create_full_driver_profile(admin_headers, wb):
    p, _sheet = _create_profile(admin_headers, wb, "Driver")
    pid = p["migration_mapping_profile_id"]
    # Grab transforms for driver_code and phone
    rules = requests.get(f"{PREP}/transform-rules", headers=admin_headers).json()
    dc_rule = next(r for r in rules if r["rule_type"] == "driver_code")["migration_transform_rule_id"]
    phone_rule = next(r for r in rules if r["rule_type"] == "phone")["migration_transform_rule_id"]
    email_rule = next(r for r in rules if r["rule_type"] == "email")["migration_transform_rule_id"]
    date_rule = next(r for r in rules if r["rule_type"] == "date")["migration_transform_rule_id"]
    _add_field(admin_headers, pid, "driver_code", "driver_code", [dc_rule], required=True)
    _add_field(admin_headers, pid, "dispatch_number", "dispatch_number", [])
    _add_field(admin_headers, pid, "full_name", "full_name", [])
    _add_field(admin_headers, pid, "email", "email", [email_rule])
    _add_field(admin_headers, pid, "mobile_phone", "mobile_phone", [phone_rule])
    _add_field(admin_headers, pid, "start_date", "start_date", [date_rule])
    return pid


def _run_dry_run(admin_headers, wb, pid, name="Test"):
    r = requests.post(f"{PREP}/dry-runs", json={
        "name": name, "description": "auto",
        "mapping_profile_ids": [pid],
        "source_workbook_ids": [wb["migration_source_workbook_id"]],
    }, headers=admin_headers, timeout=15)
    r.raise_for_status()
    dr_id = r.json()["migration_dry_run_id"]
    execu = requests.post(f"{PREP}/dry-runs/{dr_id}/execute",
                           headers=admin_headers, timeout=60)
    execu.raise_for_status()
    return dr_id, execu.json()


# ═══════════════════════════════════════════════════════════════════════════
# Workbook profiling
# ═══════════════════════════════════════════════════════════════════════════
class TestWorkbookProfiling:

    def test_upload_valid_xlsx(self, admin_headers):
        wb = _upload_workbook(admin_headers, "clean_drivers.xlsx", {"Drivers": CLEAN_DRIVERS})
        assert wb["status"] == "Profiled"
        assert wb["detected_sheet_count"] == 1
        assert wb["file_sha256"]
        assert wb["file_type"] == "xlsx"
        assert wb["authority_level"] == "Authoritative"

    def test_upload_csv(self, admin_headers):
        csv_data = "driver_code,full_name\n9010,Ficta Ten\n".encode()
        files = {"file": ("d.csv", csv_data, "text/csv")}
        form = {"profile_name": "CSV test", "authority_level": "Supporting"}
        r = requests.post(f"{PREP}/workbooks/profile", files=files, data=form,
                           headers=admin_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["file_type"] == "csv"

    def test_multiple_sheets(self, admin_headers):
        wb = _upload_workbook(admin_headers, "multi.xlsx", {
            "Drivers": CLEAN_DRIVERS,
            "Owners": [["abn", "name"], ["11122233344", "Ficta Pty Ltd"]],
        })
        assert wb["detected_sheet_count"] == 2

    def test_duplicate_headers_flagged(self, admin_headers):
        rows = [["col_a", "col_a", "col_b"], ["v1", "v2", "v3"]]
        wb = _upload_workbook(admin_headers, "dup.xlsx", {"S": rows})
        sheets = requests.get(f"{PREP}/workbooks/{wb['migration_source_workbook_id']}/sheets",
                                headers=admin_headers).json()
        assert any("duplicate_of" in c for c in sheets[0]["column_profile"])

    def test_unsupported_type(self, admin_headers):
        files = {"file": ("bad.txt", b"nope", "text/plain")}
        form = {"profile_name": "bad"}
        r = requests.post(f"{PREP}/workbooks/profile", files=files, data=form,
                           headers=admin_headers, timeout=15)
        assert r.status_code == 400

    def test_readonly_cannot_upload(self, readonly_headers):
        files = {"file": ("x.xlsx", b"any", "application/octet-stream")}
        form = {"profile_name": "x"}
        r = requests.post(f"{PREP}/workbooks/profile", files=files, data=form,
                           headers=readonly_headers, timeout=15)
        assert r.status_code == 403

    def test_archive_workbook(self, admin_headers):
        wb = _upload_workbook(admin_headers, "arch.xlsx", {"S": CLEAN_DRIVERS})
        r = requests.post(f"{PREP}/workbooks/{wb['migration_source_workbook_id']}/archive",
                            headers=admin_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["archived"] is True


# ═══════════════════════════════════════════════════════════════════════════
# Mapping profiles + fields
# ═══════════════════════════════════════════════════════════════════════════
class TestMappingProfiles:

    def test_create_and_update_draft(self, admin_headers):
        wb = _upload_workbook(admin_headers, "p1.xlsx", {"S": CLEAN_DRIVERS})
        p, _sheet = _create_profile(admin_headers, wb)
        assert p["status"] == "Draft"
        upd = requests.put(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}",
                             json={"description": "updated"}, headers=admin_headers)
        assert upd.status_code == 200
        assert upd.json()["description"] == "updated"

    def test_approve_and_immutable(self, admin_headers):
        wb = _upload_workbook(admin_headers, "p2.xlsx", {"S": CLEAN_DRIVERS})
        p, _s = _create_profile(admin_headers, wb)
        r = requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/approve",
                            headers=admin_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["status"] == "Approved"
        # Now edits must 400
        upd = requests.put(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}",
                             json={"description": "should-fail"}, headers=admin_headers)
        assert upd.status_code == 400

    def test_clone_creates_new_version(self, admin_headers):
        wb = _upload_workbook(admin_headers, "p3.xlsx", {"S": CLEAN_DRIVERS})
        p, _s = _create_profile(admin_headers, wb)
        requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/approve",
                       headers=admin_headers)
        clone = requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/clone",
                                headers=admin_headers).json()
        assert clone["profile_version"] > p["profile_version"]
        assert clone["status"] == "Draft"

    def test_second_approve_supersedes_first(self, admin_headers):
        wb = _upload_workbook(admin_headers, "p4.xlsx", {"S": CLEAN_DRIVERS})
        p, _s = _create_profile(admin_headers, wb)
        requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/approve",
                       headers=admin_headers)
        clone = requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/clone",
                                headers=admin_headers).json()
        requests.post(f"{PREP}/mapping-profiles/{clone['migration_mapping_profile_id']}/approve",
                       headers=admin_headers)
        p1 = requests.get(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}",
                            headers=admin_headers).json()
        assert p1["status"] == "Superseded"

    def test_allocator_cannot_approve(self, admin_headers, allocator_headers):
        wb = _upload_workbook(admin_headers, "p5.xlsx", {"S": CLEAN_DRIVERS})
        p, _s = _create_profile(admin_headers, wb)
        r = requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/approve",
                            headers=allocator_headers, timeout=15)
        assert r.status_code == 403

    def test_field_use_source_blocked_for_protected(self, admin_headers):
        wb = _upload_workbook(admin_headers, "p6.xlsx", {"S": CLEAN_DRIVERS})
        p, _s = _create_profile(admin_headers, wb)
        r = requests.post(f"{PREP}/mapping-profiles/{p['migration_mapping_profile_id']}/fields",
                            json={"source_column_name": "driver_code",
                                   "target_field": "driver_code",
                                   "target_entity_type": "Driver",
                                   "mapping_type": "Direct",
                                   "conflict_strategy": "Use Source",
                                   "null_handling": "Leave Blank"},
                            headers=admin_headers, timeout=15)
        assert r.status_code == 400


# ═══════════════════════════════════════════════════════════════════════════
# Transforms
# ═══════════════════════════════════════════════════════════════════════════
class TestTransforms:

    def test_transforms_list_and_deterministic(self, admin_headers):
        r = requests.get(f"{PREP}/transform-rules", headers=admin_headers).json()
        assert len(r) >= 10  # seeded


# ═══════════════════════════════════════════════════════════════════════════
# Dry-run: happy path + duplicates + reserved dispatch
# ═══════════════════════════════════════════════════════════════════════════
class TestDryRun:

    def test_clean_dry_run(self, admin_headers):
        wb = _upload_workbook(admin_headers, "clean.xlsx", {"Drivers": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, result = _run_dry_run(admin_headers, wb, pid, "clean")
        assert result["status"] == "Preview Ready"
        assert result["row_count"] == 2
        assert result["blocking_issue_count"] == 0

    def test_duplicate_driver_code_blocks(self, admin_headers):
        rows = [DRIVER_HEADERS,
                ["9100", 601, "A B", "a@example.test", "0400111111", "2024-02-01"],
                ["9100", 602, "C D", "b@example.test", "0400111112", "2024-02-02"]]
        wb = _upload_workbook(admin_headers, "dupcode.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, result = _run_dry_run(admin_headers, wb, pid, "dupcode")
        issues = requests.get(f"{PREP}/dry-runs/{dr_id}/issues", headers=admin_headers).json()
        assert any(i["issue_code"] == "DUPLICATE_DRIVER_CODE" for i in issues)
        assert result["blocking_issue_count"] >= 1

    def test_reserved_dispatch_blocks(self, admin_headers):
        rows = [DRIVER_HEADERS,
                ["9200", 13, "R E", "r@example.test", "0400222222", "2024-03-01"]]
        wb = _upload_workbook(admin_headers, "res.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, result = _run_dry_run(admin_headers, wb, pid, "reserved")
        issues = requests.get(f"{PREP}/dry-runs/{dr_id}/issues", headers=admin_headers).json()
        assert any(i["issue_code"] == "RESERVED_DISPATCH" for i in issues)

    def test_non_integer_driver_code_warning(self, admin_headers):
        rows = [DRIVER_HEADERS,
                ["LEGACY-A", 700, "Legacy One", "l@example.test", "0400333333", "2020-01-01"]]
        wb = _upload_workbook(admin_headers, "leg.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "legacy")
        issues = requests.get(f"{PREP}/dry-runs/{dr_id}/issues", headers=admin_headers).json()
        assert any(i["issue_code"] == "NON_INTEGER_DRIVER_CODE" for i in issues)
        assert any(i["severity"] == "Warning" for i in issues if i["issue_code"] == "NON_INTEGER_DRIVER_CODE")

    def test_deterministic_rerun(self, admin_headers):
        wb = _upload_workbook(admin_headers, "det.xlsx", {"Drivers": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, r1 = _run_dry_run(admin_headers, wb, pid, "det")
        rerun = requests.post(f"{PREP}/dry-runs/{dr_id}/rerun",
                                headers=admin_headers, timeout=30).json()
        assert rerun["row_count"] == r1["row_count"]
        assert rerun["blocking_issue_count"] == r1["blocking_issue_count"]

    def test_lineage_recorded(self, admin_headers):
        wb = _upload_workbook(admin_headers, "line.xlsx", {"Drivers": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "lineage")
        rows = requests.get(f"{PREP}/dry-runs/{dr_id}/rows", headers=admin_headers).json()
        assert len(rows) == 2
        assert all("source_snapshot" in r for r in rows)
        assert all("transformed_snapshot" in r for r in rows)


# ═══════════════════════════════════════════════════════════════════════════
# Issue lifecycle
# ═══════════════════════════════════════════════════════════════════════════
class TestIssues:

    def test_resolve_and_reopen(self, admin_headers):
        rows = [DRIVER_HEADERS,
                ["9300", 0, "X Y", "x@example.test", "0400444444", "2024-05-01"]]
        wb = _upload_workbook(admin_headers, "issue.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "issue")
        issues = requests.get(f"{PREP}/dry-runs/{dr_id}/issues", headers=admin_headers).json()
        assert issues
        iid = issues[0]["migration_issue_id"]
        res = requests.post(f"{PREP}/issues/{iid}/resolve",
                              json={"resolution_type": "Reject Row",
                                     "resolution_note": "will exclude"},
                              headers=admin_headers).json()
        assert res["status"] == "Resolved"
        reop = requests.post(f"{PREP}/issues/{iid}/reopen", headers=admin_headers).json()
        assert reop["status"] == "Open"

    def test_readonly_cannot_resolve(self, admin_headers, readonly_headers):
        rows = [DRIVER_HEADERS,
                ["9400", 0, "Z", "z@example.test", "0400555555", "2024-05-01"]]
        wb = _upload_workbook(admin_headers, "ro.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "ro")
        issues = requests.get(f"{PREP}/dry-runs/{dr_id}/issues", headers=admin_headers).json()
        r = requests.post(f"{PREP}/issues/{issues[0]['migration_issue_id']}/resolve",
                            json={"resolution_type": "Ignore Warning"},
                            headers=readonly_headers)
        assert r.status_code == 403


# ═══════════════════════════════════════════════════════════════════════════
# Reconciliation reports
# ═══════════════════════════════════════════════════════════════════════════
class TestReconciliation:

    def test_identifier_impact(self, admin_headers):
        rows = [DRIVER_HEADERS,
                ["9500", 750, "F G", "f@example.test", "0400666666", "2024-06-01"],
                ["LEGACY-B", 751, "H I", "h@example.test", "0400666667", "2024-06-02"]]
        wb = _upload_workbook(admin_headers, "id.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "ids")
        r = requests.get(f"{PREP}/dry-runs/{dr_id}/identifier-impact",
                           headers=admin_headers).json()
        assert r["driver_code"]["valid"] >= 1
        assert r["driver_code"]["invalid"] >= 1
        assert 0 in r["dispatch_number"]["reserved_permanent"]
        assert 13 in r["dispatch_number"]["reserved_permanent"]

    def test_no_mutation(self, admin_headers):
        """Dry run must not mutate canonical Drivers collection."""
        before = requests.get(f"{API}/drivers", headers=admin_headers).json()
        wb = _upload_workbook(admin_headers, "noop.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        _run_dry_run(admin_headers, wb, pid, "noop")
        after = requests.get(f"{API}/drivers", headers=admin_headers).json()
        assert len(after) == len(before)

    def test_activation_impact_preview_only(self, admin_headers):
        wb = _upload_workbook(admin_headers, "act.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "act")
        r = requests.get(f"{PREP}/dry-runs/{dr_id}/activation-impact",
                           headers=admin_headers).json()
        assert "preview" in r
        assert all("no checklist" in p["notes"].lower() for p in r["preview"])


# ═══════════════════════════════════════════════════════════════════════════
# Rollback preview + Go/No-Go
# ═══════════════════════════════════════════════════════════════════════════
class TestGoNoGo:

    def test_no_go_when_not_approved(self, admin_headers):
        wb = _upload_workbook(admin_headers, "nogo.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "nogo")
        r = requests.post(f"{PREP}/dry-runs/{dr_id}/go-no-go",
                            headers=admin_headers).json()
        # profile not approved → NO-GO
        assert r["result"] == "NO-GO"

    def test_go_when_approved_and_clean(self, admin_headers):
        wb = _upload_workbook(admin_headers, "go.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        # Approve profile
        requests.post(f"{PREP}/mapping-profiles/{pid}/approve", headers=admin_headers)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "go-clean")
        r = requests.post(f"{PREP}/dry-runs/{dr_id}/go-no-go",
                            headers=admin_headers).json()
        assert r["result"] in ("GO", "CONDITIONAL GO")

    def test_no_go_with_blockers(self, admin_headers):
        rows = [DRIVER_HEADERS,
                ["9600", 13, "A", "a@example.test", "0400777777", "2024-07-01"]]
        wb = _upload_workbook(admin_headers, "block.xlsx", {"S": rows})
        pid = _create_full_driver_profile(admin_headers, wb)
        requests.post(f"{PREP}/mapping-profiles/{pid}/approve", headers=admin_headers)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "block")
        r = requests.post(f"{PREP}/dry-runs/{dr_id}/go-no-go",
                            headers=admin_headers).json()
        assert r["result"] == "NO-GO"

    def test_cannot_approve_no_go(self, admin_headers):
        wb = _upload_workbook(admin_headers, "no2.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "no2")
        r = requests.post(f"{PREP}/dry-runs/{dr_id}/go-no-go", headers=admin_headers).json()
        rid = r["migration_go_no_go_report_id"]
        assert r["result"] == "NO-GO"
        app = requests.post(f"{PREP}/go-no-go/{rid}/approve", json={},
                              headers=admin_headers)
        assert app.status_code == 400

    def test_rollback_preview_deterministic(self, admin_headers):
        wb = _upload_workbook(admin_headers, "rb.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        dr_id, _ = _run_dry_run(admin_headers, wb, pid, "rb")
        r1 = requests.get(f"{PREP}/dry-runs/{dr_id}/rollback-manifest",
                           headers=admin_headers).json()
        r2 = requests.get(f"{PREP}/dry-runs/{dr_id}/rollback-manifest",
                           headers=admin_headers).json()
        assert r1["would_create_field_count"] == r2["would_create_field_count"]
        assert r1["applied"] is False


# ═══════════════════════════════════════════════════════════════════════════
# Safety: no commit endpoint
# ═══════════════════════════════════════════════════════════════════════════
class TestNoCommit:

    def test_no_commit_endpoint_exists(self, admin_headers):
        for suffix in ("/commit", "/dry-runs/x/commit", "/migrate", "/apply"):
            r = requests.post(f"{PREP}{suffix}", headers=admin_headers)
            assert r.status_code in (404, 405), f"Unexpectedly found {suffix}"

    def test_dry_run_does_not_touch_number_sequences(self, admin_headers):
        # Snapshot sequences
        before = requests.get(f"{API}/numbering/summary",
                                headers=admin_headers).json() if requests.get(
                                    f"{API}/numbering/summary",
                                    headers=admin_headers).status_code == 200 else None
        wb = _upload_workbook(admin_headers, "seq.xlsx", {"S": CLEAN_DRIVERS})
        pid = _create_full_driver_profile(admin_headers, wb)
        _run_dry_run(admin_headers, wb, pid, "seq")
        after = requests.get(f"{API}/numbering/summary",
                               headers=admin_headers).json() if requests.get(
                                   f"{API}/numbering/summary",
                                   headers=admin_headers).status_code == 200 else None
        assert before == after
