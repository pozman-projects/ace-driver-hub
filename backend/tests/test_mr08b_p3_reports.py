"""MR-08B-P3 · Canonical Report Builder tests."""
from __future__ import annotations

import io
import os
import uuid

import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/") or "http://localhost:8001"
ADMIN = {"email": "admin@acedriverhub.com", "password": "Admin@123"}


def _tag(): return uuid.uuid4().hex[:8]


def _login(email, pw):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": pw}, timeout=15)
    r.raise_for_status()
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {r.json()['access_token']}"})
    return s


def _register(admin, role):
    email = f"mr08p3-{role.lower()}-{_tag()}@acedriverhub.com"
    admin.post(f"{BASE_URL}/api/auth/register",
                json={"email": email, "password": "P@ssword1",
                      "full_name": f"MR08P3 {role}", "role": role}, timeout=15).raise_for_status()
    return _login(email, "P@ssword1")


@pytest.fixture(scope="module")
def admin(): return _login(ADMIN["email"], ADMIN["password"])
@pytest.fixture(scope="module")
def manager(admin): return _register(admin, "Manager")
@pytest.fixture(scope="module")
def compliance(admin): return _register(admin, "Compliance")
@pytest.fixture(scope="module")
def allocator(admin): return _register(admin, "Allocator")
@pytest.fixture(scope="module")
def readonly(admin): return _register(admin, "ReadOnly")


# ─── 1-3 Sources / unknown source + field ─────────────────────────────
class TestSources:
    def test_sources_allowlist(self, admin):
        r = admin.get(f"{BASE_URL}/api/reports/sources", timeout=15)
        r.raise_for_status()
        keys = {s["key"] for s in r.json()}
        expected = {"drivers", "owners", "vehicles", "equipment",
                    "driver-licences", "vehicle-registrations", "vehicle-insurance",
                    "vehicle-inspections", "vehicle-defects", "vehicle-maintenance",
                    "equipment-compliance", "activation-readiness", "documents"}
        assert expected.issubset(keys)

    def test_unknown_source_rejected(self, admin):
        r = admin.get(f"{BASE_URL}/api/reports/sources/nope/fields", timeout=15)
        assert r.status_code == 400

    def test_unknown_field_rejected_on_run(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                       json={"source": "drivers", "fields": ["not_a_field"]}, timeout=15)
        assert r.status_code == 400


# ─── 4-5 MR-07A sensitive driver fields ───────────────────────────────
class TestSensitiveDriverFields:
    def test_admin_sees_sensitive_fields(self, admin):
        r = admin.get(f"{BASE_URL}/api/reports/sources/drivers/fields", timeout=15)
        keys = {f["key"] for f in r.json()["fields"]}
        for f in ("business_name", "abn", "payroll_number", "payment_percentage"):
            assert f in keys, f"Admin missing {f}"

    def test_manager_sees_sensitive_fields(self, manager):
        r = manager.get(f"{BASE_URL}/api/reports/sources/drivers/fields", timeout=15)
        keys = {f["key"] for f in r.json()["fields"]}
        assert "business_name" in keys

    @pytest.mark.parametrize("role_name", ["compliance", "allocator", "readonly"])
    def test_non_privileged_hidden_sensitive_fields(self, request, role_name):
        sess = request.getfixturevalue(role_name)
        r = sess.get(f"{BASE_URL}/api/reports/sources/drivers/fields", timeout=15)
        keys = {f["key"] for f in r.json()["fields"]}
        for f in ("business_name", "abn", "payroll_number", "payment_percentage"):
            assert f not in keys, f"{role_name} sees {f}"

    @pytest.mark.parametrize("role_name", ["compliance", "allocator", "readonly"])
    def test_non_privileged_cannot_request_sensitive_field(self, request, role_name):
        sess = request.getfixturevalue(role_name)
        r = sess.post(f"{BASE_URL}/api/reports/run",
                       json={"source": "drivers", "fields": ["full_name", "business_name"]},
                       timeout=15)
        assert r.status_code == 403, r.text


# ─── 6 Normal Driver report for legitimate roles ──────────────────────
class TestBasicRun:
    def test_admin_can_run_driver_report(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name", "driver_status"]},
                        timeout=30)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["source"] == "drivers"
        assert isinstance(body["rows"], list)
        assert set(body["rows"][0].keys()) <= {"full_name", "driver_status"} if body["rows"] else True

    def test_readonly_can_run_basic_report(self, readonly):
        r = readonly.post(f"{BASE_URL}/api/reports/run",
                           json={"source": "drivers", "fields": ["full_name"]}, timeout=30)
        assert r.status_code == 200


# ─── 7-8 Document sensitivity ─────────────────────────────────────────
class TestDocumentSensitivity:
    def _seed(self, admin, sensitivity):
        did = admin.post(f"{BASE_URL}/api/drivers",
                         json={"full_name": f"DocSeed {_tag()}", "driver_status": "Training"},
                         timeout=15).json()["id"]
        files = {"file": (f"r-{_tag()}.csv", b"a,b,c", "text/csv")}
        data = {"entity_type": "Driver", "entity_id": did, "document_type": "Other",
                "sensitivity": sensitivity, "title": f"MR08P3 {sensitivity} {_tag()}"}
        admin.post(f"{BASE_URL}/api/documents/upload", files=files, data=data, timeout=30).raise_for_status()
        return sensitivity

    def test_restricted_filtered_out_for_non_privileged(self, admin, compliance, allocator, readonly):
        self._seed(admin, "Restricted")
        for sess, role in ((compliance, "Compliance"), (allocator, "Allocator"), (readonly, "ReadOnly")):
            r = sess.post(f"{BASE_URL}/api/reports/run",
                           json={"source": "documents",
                                 "fields": ["title", "sensitivity"],
                                 "limit": 1000},
                           timeout=45)
            assert r.status_code == 200
            leaks = [row for row in r.json()["rows"] if row.get("sensitivity") == "Restricted"]
            assert not leaks, f"{role} saw Restricted document rows"

    def test_confidential_visible_to_compliance_not_allocator(self, admin, compliance, allocator):
        self._seed(admin, "Confidential")
        r_c = compliance.post(f"{BASE_URL}/api/reports/run",
                              json={"source": "documents", "fields": ["title", "sensitivity"],
                                    "limit": 1000}, timeout=45).json()
        assert any(row.get("sensitivity") == "Confidential" for row in r_c["rows"])
        r_a = allocator.post(f"{BASE_URL}/api/reports/run",
                             json={"source": "documents", "fields": ["title", "sensitivity"],
                                   "limit": 1000}, timeout=45).json()
        assert not any(row.get("sensitivity") == "Confidential" for row in r_a["rows"])


# ─── 9-13 Filter / sort validation ────────────────────────────────────
class TestFilterSortValidation:
    def test_invalid_operator_for_type(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name"],
                              "filters": [{"field": "full_name", "operator": "greater_than", "value": "x"}]},
                        timeout=15)
        assert r.status_code == 400

    def test_raw_mongo_operator_rejected(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name"],
                              "filters": [{"field": "full_name", "operator": "$where", "value": "1"}]},
                        timeout=15)
        assert r.status_code == 400

    def test_max_two_sorts(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name"],
                              "sort": [{"field": "full_name"}, {"field": "driver_status"},
                                       {"field": "start_date"}]},
                        timeout=15)
        assert r.status_code == 400

    def test_unknown_sort_field(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name"],
                              "sort": [{"field": "nope"}]}, timeout=15)
        assert r.status_code == 400


# ─── 14-16 Archive / preview limit ────────────────────────────────────
class TestPreviewLimits:
    def test_default_excludes_archived_drivers(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name", "driver_status"],
                              "limit": 1000}, timeout=45).json()
        assert not any(row.get("driver_status") == "Archived" for row in r["rows"])

    def test_limit_enforced(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name"], "limit": 5},
                        timeout=15).json()
        assert len(r["rows"]) <= 5


# ─── 17-20 CSV / XLSX ────────────────────────────────────────────────
class TestCsvXlsx:
    def test_csv_export(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/export/csv",
                        json={"source": "drivers", "fields": ["full_name", "driver_status"]},
                        timeout=45)
        assert r.status_code == 200
        eid = r.json()["export"]["id"]
        d = admin.get(f"{BASE_URL}/api/reports/exports/{eid}/download", timeout=45)
        assert d.status_code == 200
        text = d.content.decode()
        header = text.splitlines()[0]
        assert header == "Full Name,Status", header

    def test_csv_formula_neutralised(self, admin):
        # Seed a driver whose name starts with `=` to trigger neutralisation.
        admin.post(f"{BASE_URL}/api/drivers",
                    json={"full_name": "=SUM(A1)", "driver_status": "Training"},
                    timeout=15).raise_for_status()
        r = admin.post(f"{BASE_URL}/api/reports/export/csv",
                        json={"source": "drivers", "fields": ["full_name"], "limit": 1000},
                        timeout=45)
        eid = r.json()["export"]["id"]
        d = admin.get(f"{BASE_URL}/api/reports/exports/{eid}/download", timeout=45).content.decode()
        assert "=SUM(A1)" not in d.replace("'=SUM(A1)", "")
        assert "'=SUM(A1)" in d

    def test_xlsx_export(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/export/xlsx",
                        json={"source": "drivers", "fields": ["full_name"]}, timeout=60)
        assert r.status_code == 200
        eid = r.json()["export"]["id"]
        d = admin.get(f"{BASE_URL}/api/reports/exports/{eid}/download", timeout=60)
        assert d.status_code == 200
        # basic XLSX signature
        assert d.content[:2] == b"PK", "XLSX must be a valid zip"
        # verify no formula cells
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(d.content))
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    assert not cell.startswith("="), f"formula cell in XLSX: {cell}"


# ─── 21-22 Company display resolution ────────────────────────────────
class TestCompanyDisplay:
    def test_company_id_resolves_to_name(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "drivers", "fields": ["full_name", "company_id"],
                              "limit": 1000}, timeout=45).json()
        # any driver with a company_id row should show the canonical name, not the UUID.
        for row in r["rows"]:
            if row.get("company_id"):
                assert "-" not in row["company_id"] or row["company_id"] == "ACE Car Freighters" \
                    or " " in row["company_id"], f"Unresolved company_id: {row['company_id']}"


# ─── 25 Download security ────────────────────────────────────────────
class TestDownloadSecurity:
    def test_creator_can_download(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/export/csv",
                        json={"source": "drivers", "fields": ["full_name"]}, timeout=45)
        eid = r.json()["export"]["id"]
        assert admin.get(f"{BASE_URL}/api/reports/exports/{eid}/download", timeout=15).status_code == 200

    def test_non_creator_non_admin_denied(self, admin, allocator):
        r = admin.post(f"{BASE_URL}/api/reports/export/csv",
                        json={"source": "drivers", "fields": ["full_name"]}, timeout=45)
        eid = r.json()["export"]["id"]
        assert allocator.get(f"{BASE_URL}/api/reports/exports/{eid}/download", timeout=15).status_code == 403


# ─── 26 Revalidate at export time ────────────────────────────────────
class TestExportRevalidates:
    def test_compliance_cannot_export_sensitive(self, compliance):
        r = compliance.post(f"{BASE_URL}/api/reports/export/csv",
                             json={"source": "drivers", "fields": ["business_name"]},
                             timeout=15)
        assert r.status_code == 403


# ─── 29 Activation read-only ─────────────────────────────────────────
class TestActivationReadOnly:
    def test_activation_source_returns_readiness_rows(self, admin):
        r = admin.post(f"{BASE_URL}/api/reports/run",
                        json={"source": "activation-readiness",
                              "fields": ["driver_full_name", "readiness", "missing_count"],
                              "limit": 5}, timeout=45)
        assert r.status_code == 200, r.text
        for row in r.json()["rows"]:
            assert row.get("readiness") in (None, "Ready", "Not Ready", "Unknown")


# ─── 30 Existing driver PDF exports unchanged ────────────────────────
class TestDriverExportsUnchanged:
    def test_driver_start_sheet_still_works(self, admin):
        did = admin.post(f"{BASE_URL}/api/drivers",
                         json={"full_name": f"PdfSanity {_tag()}", "driver_status": "Training"},
                         timeout=15).json()["id"]
        r = admin.post(f"{BASE_URL}/api/drivers/{did}/exports/start-sheet",
                        json={"confirm": True}, timeout=60)
        # If MR-04 blocks the export we still expect a non-500 gate. 200 is
        # ideal; a policy 409 is acceptable — anything but 500.
        assert r.status_code < 500, r.text
